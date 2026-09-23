import os
import docx

def create_docx_sample():
    os.makedirs("samples", exist_ok=True)
    doc = docx.Document()
    doc.add_heading("NON-DISCLOSURE AND INTELLECTUAL PROPERTY AGREEMENT", 0)
    
    p = doc.add_paragraph()
    p.add_run("This Agreement is made on October 1, 2026, between Alpha Corp and Beta Labs.")
    
    doc.add_heading("1. Confidentiality Obligations", level=1)
    doc.add_paragraph("Receiving Party agrees to protect Confidential Information with reasonable care.")
    
    doc.add_heading("2. Intellectual Property Ownership", level=1)
    doc.add_paragraph("All work product and inventions created shall be irrevocably assigned to Alpha Corp.")
    
    doc.add_heading("3. Limitation of Liability", level=1)
    doc.add_paragraph("Neither party shall be liable for indirect damages. Aggregate liability is capped at total fees paid.")

    doc.add_heading("4. Governing Law", level=1)
    doc.add_paragraph("This agreement shall be governed by the laws of California.")

    doc.save("samples/Sample_NDA_Agreement.docx")
    print("Created samples/Sample_NDA_Agreement.docx")

def create_txt_sample():
    os.makedirs("samples", exist_ok=True)
    with open("samples/Enterprise_SaaS_Agreement.txt", "w", encoding="utf-8") as f:
        f.write("""MASTER SAAS AGREEMENT

1. SUBSCRIPTION SERVICES
Provider grants Customer access to cloud services subject to timely fee payments.

2. INDEMNIFICATION AND LIABILITIES
Customer shall indemnify Provider against third-party claims. Provider offers no indemnification. Customer liability is unlimited.

3. TERMINATION AND CURE
Provider may terminate immediately without notice upon any breach. Customer may terminate only upon 60 days written notice.

4. RESTRICTIVE COVENANTS
Customer shall not solicit or hire Provider employees for 24 months post-termination.
""")
    print("Created samples/Enterprise_SaaS_Agreement.txt")

if __name__ == "__main__":
    create_docx_sample()
    create_txt_sample()
