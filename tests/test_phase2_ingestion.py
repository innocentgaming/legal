import io
import os
import pytest
import docx
from fastapi.testclient import TestClient
from backend.main import app
from backend.app.services.document_parser import DocumentParserService
from backend.app.services.clause_segmentation_service import ClauseSegmentationService
from backend.app.services.ingestion_service import IngestionService
from backend.app.core.errors import DocumentProcessingError

client = TestClient(app)

def create_sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", 0)
    
    doc.add_heading("1. Confidential Information", level=1)
    doc.add_paragraph("1.1 Definition. Confidential Information includes all proprietary technical and commercial data disclosed by Disclosing Party.")
    doc.add_paragraph("1.2 Standard of Care. Receiving Party shall maintain confidentiality using at least reasonable care.")
    
    doc.add_heading("2. Term and Termination", level=1)
    doc.add_paragraph("2.1 Duration. This Agreement shall remain in effect for two (2) years from the Effective Date.")
    doc.add_paragraph("2.2 Survival. Confidentiality obligations shall survive termination for five (5) years.")

    doc.add_heading("3. Governing Law", level=1)
    doc.add_paragraph("3.1 Jurisdiction. This Agreement is governed by the laws of the State of Delaware.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_unit_parser_and_segmentation():
    sample_text = """MASTER SERVICES AGREEMENT

1. SERVICES AND ACCESS
1.1 Subscription. Provider shall provide access to cloud services.
1.2 Restrictions. Customer shall not reverse engineer the platform.

2. INDEMNIFICATION AND LIABILITY
2.1 Indemnity. Customer shall indemnify Provider against third-party claims.
2.2 Liability Cap. In no event shall Provider liability exceed fees paid in the prior 12 months.

3. GOVERNING LAW
This agreement is governed by the laws of New York.
"""
    # 1. Parsing
    parsed = DocumentParserService.parse("contract.txt", sample_text.encode("utf-8"))
    assert parsed["document_type"] == "txt"
    assert parsed["page_count"] >= 1
    
    # 2. Structural Segmentation
    sections = ClauseSegmentationService.segment_document(parsed)
    assert len(sections) >= 3
    assert sections[0].section_id == "SEC-001"
    assert sections[0].clauses[0].clause_id == "CLAUSE-001"
    assert "SERVICES" in sections[0].title.upper()

    # Verify stable IDs and numbering
    all_clauses = [c for s in sections for c in s.clauses]
    assert all_clauses[0].clause_id == "CLAUSE-001"
    assert all_clauses[1].clause_id == "CLAUSE-002"
    assert all_clauses[2].clause_id == "CLAUSE-003"

def test_scenario_1_pdf_upload():
    # Scenario 1: PDF ingestion with pdfplumber
    with open("samples/sample_contract.pdf", "rb") as pdf_file:
        res = client.post(
            "/api/documents/upload",
            files={"file": ("sample_contract.pdf", pdf_file.read(), "application/pdf")}
        )
    assert res.status_code == 200, f"Upload PDF failed: {res.text}"
    data = res.json()
    assert data["document_type"] == "pdf"
    assert data["page_count"] >= 1
    assert data["clause_count"] >= 1
    assert data["processing_status"] == "completed"
    assert len(data["sections"]) >= 1
    assert data["sections"][0]["clauses"][0]["clause_id"] == "CLAUSE-001"
    assert "metadata" in data
    assert "page_count" in data["metadata"]

def test_scenario_2_docx_upload():
    # Scenario 2: DOCX ingestion with Mammoth
    docx_bytes = create_sample_docx_bytes()
    res = client.post(
        "/api/documents/upload",
        files={"file": ("sample_nda.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert res.status_code == 200, f"Upload DOCX failed: {res.text}"
    data = res.json()
    assert data["document_type"] == "docx"
    assert data["clause_count"] >= 3
    assert data["processing_status"] == "completed"
    assert len(data["sections"]) >= 3
    assert data["sections"][0]["clauses"][0]["clause_id"] == "CLAUSE-001"

def test_scenario_3_txt_upload():
    # Scenario 3: TXT native processing
    txt_content = """SOFTWARE LICENSE AGREEMENT

1. GRANT OF LICENSE
1.1 License Scope. Licensor grants Licensee a non-exclusive license to use the Software.
1.2 Restrictions. Licensee shall not sublicense or modify the binary code.

2. PAYMENT TERMS
2.1 Annual Fee. Licensee shall pay $12,000 annually.
"""
    res = client.post(
        "/api/documents/upload",
        files={"file": ("license_agreement.txt", txt_content.encode("utf-8"), "text/plain")}
    )
    assert res.status_code == 200, f"Upload TXT failed: {res.text}"
    data = res.json()
    assert data["document_type"] == "txt"
    assert data["clause_count"] >= 2
    assert data["processing_status"] == "completed"
    assert data["sections"][0]["clauses"][0]["clause_id"] == "CLAUSE-001"

def test_scenario_4_empty_document_rejection():
    # Scenario 4: Empty document rejected with 400 Bad Request
    res = client.post(
        "/api/documents/upload",
        files={"file": ("empty.txt", b"", "text/plain")}
    )
    assert res.status_code == 400
    assert "empty" in res.json().get("detail", "").lower() or "empty" in str(res.json()).lower()

def test_scenario_5_invalid_file_type_rejection():
    # Scenario 5: Unsupported format rejected
    res = client.post(
        "/api/documents/upload",
        files={"file": ("malicious.exe", b"MZ\x90\x00ExecutableData", "application/x-msdownload")}
    )
    assert res.status_code == 400
    assert "unsupported" in res.json().get("detail", "").lower() or "unsupported" in str(res.json()).lower()

def test_scenario_6_large_file_rejection():
    # Scenario 6: File exceeding MAX_FILE_SIZE_MB (25MB) rejected
    huge_payload = b"A" * (26 * 1024 * 1024)
    res = client.post(
        "/api/documents/upload",
        files={"file": ("huge_file.txt", huge_payload, "text/plain")}
    )
    assert res.status_code in [400, 413, 422]

def test_scenario_7_numbered_clauses_preservation():
    # Scenario 7: Document with numbered clauses preserves numbering and generates stable IDs
    contract = """CLOUD SERVICES AGREEMENT

1. PROVISION OF SERVICES
1.1 Service Level. Provider guarantees 99.9% uptime.
1.2 Scheduled Maintenance. Provider will notify Customer 48 hours prior.

2. FEES AND BILLING
2.1 Invoicing. Invoices are issued monthly and payable net-30.
2.2 Taxes. Customer is responsible for all applicable sales taxes.

3. CONFIDENTIALITY
3.1 Non-Disclosure. Both parties agree to keep proprietary data strictly confidential.
"""
    res = client.post(
        "/api/documents/upload",
        files={"file": ("numbered_contract.txt", contract.encode("utf-8"), "text/plain")}
    )
    assert res.status_code == 200
    data = res.json()
    
    # Verify clause numbers and stable IDs
    clauses = data["clauses"]
    clause_ids = [c["clause_id"] for c in clauses]
    assert "CLAUSE-001" in clause_ids
    assert "CLAUSE-002" in clause_ids
    assert "CLAUSE-003" in clause_ids
    
    # Verify clause numbering preservation
    clause_numbers = [c["clause_number"] for c in clauses]
    assert any("1.1" in cn for cn in clause_numbers)
    assert any("2.1" in cn for cn in clause_numbers)
    assert any("3.1" in cn for cn in clause_numbers)

def test_scenario_8_headings_and_structure_preservation():
    # Scenario 8: Document with headings preserves section hierarchy and metadata
    contract = """ENTERPRISE VENDOR AGREEMENT

SECTION 1. DEFINITIONS
"Confidential Information" shall have the broadest legal definition.

SECTION 2. INDEMNIFICATION
Vendor shall defend and indemnify Customer against any patent infringement claims.

SECTION 3. LIMITATION OF LIABILITY
Neither party shall be liable for lost profits or punitive damages.

SECTION 4. TERMINATION
Either party may terminate for cause upon thirty (30) days written notice.
"""
    res = client.post(
        "/api/documents/upload",
        files={"file": ("vendor_agreement.txt", contract.encode("utf-8"), "text/plain")}
    )
    assert res.status_code == 200
    data = res.json()
    
    # Verify section titles
    section_titles = [s["title"].upper() for s in data["sections"]]
    assert any("DEFINITIONS" in st for st in section_titles)
    assert any("INDEMNIFICATION" in st for st in section_titles)
    assert any("LIMITATION OF LIABILITY" in st for st in section_titles)
    assert any("TERMINATION" in st for st in section_titles)
    
    # Verify internal schema tree structure
    assert "document_id" in data
    assert "filename" in data
    assert "document_type" in data
    assert "page_count" in data
    assert "clause_count" in data
    assert "processing_status" in data
    assert "metadata" in data
    assert "sections" in data
    assert "clauses" in data
    
    for section in data["sections"]:
        assert "section_id" in section
        assert "title" in section
        assert "clauses" in section
        for cl in section["clauses"]:
            assert "clause_id" in cl
            assert "clause_number" in cl
            assert "original_text" in cl
            assert "page" in cl
            assert "metadata" in cl

def test_real_sample_contract_extraction():
    # Upload and verify sample contract
    with open("samples/Enterprise_SaaS_Agreement.txt", "rb") as f:
        res = client.post(
            "/api/documents/upload",
            files={"file": ("Enterprise_SaaS_Agreement.txt", f.read(), "text/plain")}
        )
    assert res.status_code == 200
    data = res.json()
    assert data["processing_status"] == "completed"
    assert data["clause_count"] >= 4
    assert len(data["sections"]) >= 4
    assert data["clauses"][0]["clause_id"] == "CLAUSE-001"
    
    # Also verify current document session retrieval
    res_curr = client.get("/api/documents/current")
    assert res_curr.status_code == 200
    curr_data = res_curr.json()
    assert curr_data["document_id"] == data["document_id"]

if __name__ == "__main__":
    pytest.main(["-v", __file__])
